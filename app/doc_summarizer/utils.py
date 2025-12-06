from docx import Document
from io import BytesIO



def extract_docx_text(content: bytes) -> str:
    doc = Document(BytesIO(content))
    text_lines = []

    # Extract paragraphs
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():  # ignore empty lines
            text_lines.append(paragraph.text)

    # Extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join(cell.text for cell in row.cells if cell.text.strip())
            if row_text:
                text_lines.append(row_text)

    return "\n".join(text_lines)